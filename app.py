import io
import re
import time
import streamlit as st
from google import genai
from google.genai import types

# ---------- Page setup ----------
st.set_page_config(page_title="Gemini Flash Studio", page_icon="✨", layout="centered")

st.markdown(
    """
    <style>
    .stApp {
        background: linear-gradient(160deg, #F7F9FF 0%, #EEF2FF 45%, #FDF4FF 100%);
    }
    .hero {
        text-align: center;
        padding: 1.6rem 1rem 1rem 1rem;
    }
    .hero h1 {
        font-size: 2.4rem;
        margin-bottom: 0.2rem;
        background: linear-gradient(90deg, #6C63FF, #EC6FD2);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }
    .hero p { color: #6B7280; font-size: 1.05rem; }
    [data-testid="stSidebar"] {
        background: #FFFFFF;
        border-right: 1px solid #E5E7EB;
    }
    .stButton > button {
        background: linear-gradient(90deg, #6C63FF, #8B7CFF);
        color: white;
        border: none;
        border-radius: 12px;
        padding: 0.55rem 1.4rem;
        font-weight: 600;
    }
    .stButton > button:hover { filter: brightness(1.07); color: white; }
    [data-testid="stChatMessage"] {
        background: #FFFFFF;
        border: 1px solid #E8EAF6;
        border-radius: 16px;
        box-shadow: 0 2px 8px rgba(108, 99, 255, 0.06);
    }
    textarea { border-radius: 12px !important; }

    /* ---- Force readable dark text (works even if Streamlit is in dark mode) ---- */
    .stApp, .stApp p, .stApp li, .stApp label, .stApp span,
    .stApp h2, .stApp h3, .stApp h4, .stApp code,
    [data-testid="stMarkdownContainer"] *,
    [data-testid="stChatMessage"] *,
    [data-testid="stSidebar"] * { color: #1F2937 !important; }

    .hero p { color: #6B7280 !important; }
    .hero h1 { -webkit-text-fill-color: transparent !important; }

    /* Tabs */
    .stTabs [data-baseweb="tab"] p { color: #4B5563 !important; font-weight: 600; }
    .stTabs [aria-selected="true"] p { color: #6C63FF !important; }
    .stTabs [data-baseweb="tab-highlight"] { background-color: #6C63FF !important; }

    /* Inputs: chat box, text area, text input, dropdowns */
    [data-testid="stChatInput"] { background: #FFFFFF !important; border: 1px solid #D9DDF5; border-radius: 16px; }
    [data-testid="stChatInput"] textarea { background: #FFFFFF !important; color: #1F2937 !important; }
    [data-testid="stChatInput"] textarea::placeholder { color: #9CA3AF !important; }
    [data-testid="stBottom"], [data-testid="stBottom"] > div { background: transparent !important; }
    textarea, input { background: #FFFFFF !important; color: #1F2937 !important; }
    [data-baseweb="select"] > div { background: #FFFFFF !important; }
    [data-baseweb="select"] * { color: #1F2937 !important; }
    [data-testid="stHeader"] { background: transparent !important; }

    /* Chat input: remove dark outer box */
    [data-testid="stChatInput"] > div,
    [data-testid="stChatInput"] div[data-baseweb="textarea"],
    [data-testid="stChatInput"] div[data-baseweb="base-input"] { background: #FFFFFF !important; }
    [data-testid="stChatInputSubmitButton"] { background: #6C63FF !important; border-radius: 10px; }
    [data-testid="stChatInputSubmitButton"] * { color: #FFFFFF !important; }

    /* Dropdown menus */
    [data-baseweb="popover"], [data-baseweb="popover"] > div,
    [data-baseweb="menu"], [data-baseweb="menu"] ul { background: #FFFFFF !important; }
    [data-baseweb="menu"] li, [data-baseweb="menu"] li * { color: #1F2937 !important; background: transparent !important; }
    [data-baseweb="menu"] li:hover, [data-baseweb="menu"] li[aria-selected="true"] { background: #EEF2FF !important; }

    /* "Thinking" indicator */
    .thinking {
        display: inline-flex; align-items: center; gap: 12px;
        padding: 10px 18px; border-radius: 999px;
        background: #F3F1FF; border: 1px solid #E0DCFF;
        font-weight: 600;
        animation: pulse-glow 2s ease-in-out infinite;
    }
    .thinking, .thinking * { color: #5B52E0 !important; }
    .thinking .dots span {
        display: inline-block; width: 8px; height: 8px; margin: 0 2px;
        border-radius: 50%; background: #6C63FF;
        animation: bounce-dot 1.2s infinite ease-in-out;
    }
    .thinking .dots span:nth-child(2) { animation-delay: 0.2s; }
    .thinking .dots span:nth-child(3) { animation-delay: 0.4s; }
    @keyframes bounce-dot {
        0%, 80%, 100% { transform: scale(0.6); opacity: 0.4; }
        40% { transform: scale(1); opacity: 1; }
    }
    @keyframes pulse-glow {
        0%, 100% { box-shadow: 0 0 0 0 rgba(108, 99, 255, 0.25); }
        50% { box-shadow: 0 0 0 8px rgba(108, 99, 255, 0); }
    }

    /* File uploader */
    [data-testid="stFileUploaderDropzone"] {
        background: #FFFFFF !important; border: 2px dashed #C9C5FF !important; border-radius: 16px !important;
    }
    [data-testid="stFileUploaderFile"] { background: #F3F1FF !important; border-radius: 12px; }
    [data-testid="stFileUploaderDropzone"] button {
        background: linear-gradient(90deg, #6C63FF, #8B7CFF) !important;
        border: none !important; border-radius: 12px !important;
    }
    [data-testid="stFileUploaderDropzone"] button, [data-testid="stFileUploaderDropzone"] button * { color: #FFFFFF !important; }

    /* Buttons keep white text (must come last) */
    .stButton > button, .stButton > button * { color: #FFFFFF !important; }
    </style>
    <div class="hero">
        <h1>✨ Gemini Flash Studio</h1>
        <p>Chat, summarize and polish your writing with Google's fast Gemini model.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

MODEL_OPTIONS = ["gemini-3.8-flash", "gemini-3.7-flash", "gemini-3.5-flash-lite"]

# ---------- Sidebar ----------
with st.sidebar:
    st.header("⚙️ Settings")

    api_key = ""
    try:
        api_key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        pass
    if not api_key:
        api_key = st.text_input("Gemini API key", type="password",
                                help="Get a free key at aistudio.google.com")

    chosen = st.selectbox("Model", MODEL_OPTIONS)
    custom = st.text_input("Custom model name (optional)", placeholder="e.g. gemini-3.8-flash",
                           help="Leave empty to use the model selected above.")
    model = custom.strip() or chosen
    temperature = st.slider("Creativity", 0.0, 2.0, 1.0, 0.1)

    if st.button("🗑️ Clear chat"):
        st.session_state.messages = []
        st.rerun()

    st.caption("Your key is used only in this session and is never stored.")

if not api_key:
    st.info("👈 Add your Gemini API key in the sidebar to get started.")
    st.stop()

client = genai.Client(api_key=api_key)


def wait_seconds(msg):
    """Read 'retry in 18s' style hints out of a Google error message."""
    m = re.search(r"retry in (\d+(?:\.\d+)?)s", msg) or re.search(r"'retryDelay': '(\d+)s'", msg)
    return int(float(m.group(1))) + 1 if m else None


def stream_reply(contents, system):
    """Stream a reply. Retries when Google is overloaded (503) and
    switches to another model when the free quota (429) is used up."""
    config = types.GenerateContentConfig(system_instruction=system, temperature=temperature)
    models_to_try = [model] + [m for m in MODEL_OPTIONS if m != model]
    last_error = None

    for m in models_to_try:
        for attempt in range(3):
            got_text = False
            try:
                for chunk in client.models.generate_content_stream(
                    model=m, contents=contents, config=config
                ):
                    if chunk.text:
                        got_text = True
                        yield chunk.text
                return
            except Exception as e:
                last_error = e
                text = str(e)
                quota = "429" in text or "RESOURCE_EXHAUSTED" in text
                overloaded = "503" in text or "UNAVAILABLE" in text
                missing = "404" in text or "NOT_FOUND" in text
                if got_text or not (quota or overloaded or missing):
                    raise
                if quota or missing:
                    break  # quota used up or model retired -> try the next model
                time.sleep(2 ** attempt)  # overloaded: wait 1s, 2s, 4s and retry
    raise last_error


def show_error(e):
    msg = str(e)
    if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
        secs = wait_seconds(msg)
        wait = f"about {secs} seconds" if secs else "a minute"
        st.warning(
            f"⏳ Free-tier limit reached (only a few requests per minute are allowed). "
            f"Please wait {wait} and try again. If it keeps happening, you may have used "
            f"up today's free quota. It resets daily."
        )
    elif "404" in msg or "NOT_FOUND" in msg:
        st.error("This model name isn't available any more. Pick another model in the sidebar.")
    elif "503" in msg or "UNAVAILABLE" in msg:
        st.warning("Google's servers are very busy right now. Please wait a few seconds and try again.")
    elif any(k in msg for k in ("API key", "API_KEY", "401", "403")):
        st.error("Your API key looks invalid. Please check it in the sidebar.")
    else:
        st.error(f"Something went wrong: {e}")


THINKING_HTML = (
    '<div class="thinking"><span>{msg}</span>'
    '<span class="dots"><span></span><span></span><span></span></span></div>'
)


def run_with_indicator(contents, system, message="✨ Gemini is thinking"):
    """Show an animated 'working' pill until the first words arrive, then stream the answer."""
    holder = st.empty()
    holder.markdown(THINKING_HTML.replace("{msg}", message), unsafe_allow_html=True)

    def gen():
        first = True
        try:
            for chunk in stream_reply(contents, system):
                if first:
                    holder.empty()
                    first = False
                yield chunk
        finally:
            holder.empty()

    return st.write_stream(gen())


FILE_TYPES = ["pdf", "docx", "txt", "md", "csv", "png", "jpg", "jpeg"]
IMAGE_MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg"}
MAX_MB = 15


def prepare_file(f):
    """Turn an uploaded file into model input. Returns (contents, error_message)."""
    data = f.getvalue()
    if len(data) > MAX_MB * 1024 * 1024:
        return None, f"This file is too large. Please keep it under {MAX_MB} MB."
    ext = f.name.rsplit(".", 1)[-1].lower()
    try:
        if ext == "pdf":  # Gemini reads PDFs (even scanned ones) directly
            return [types.Part.from_bytes(data=data, mime_type="application/pdf"), "Here is the document."], None
        if ext in IMAGE_MIME:  # Gemini can read text and content in images
            return [types.Part.from_bytes(data=data, mime_type=IMAGE_MIME[ext]), "Here is the image."], None
        if ext == "docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(lines)
        else:  # txt, md, csv
            text = data.decode("utf-8", errors="ignore")
    except Exception as e:
        return None, f"Couldn't read this file: {e}"
    if not text.strip():
        return None, "No readable text was found in this file."
    return text, None


# ---------- Tabs ----------
chat_tab, tools_tab = st.tabs(["💬 Chat", "🛠️ Writing Tools"])

# ----- Chat -----
with chat_tab:
    if "messages" not in st.session_state:
        st.session_state.messages = []

    for m in st.session_state.messages:
        with st.chat_message(m["role"]):
            st.markdown(m["content"])

    if prompt := st.chat_input("Ask me anything..."):
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        history = [
            types.Content(
                role="user" if m["role"] == "user" else "model",
                parts=[types.Part.from_text(text=m["content"])],
            )
            for m in st.session_state.messages
        ]

        with st.chat_message("assistant"):
            try:
                reply = run_with_indicator(
                    history, "You are a friendly, clear and helpful assistant."
                )
                st.session_state.messages.append({"role": "assistant", "content": reply})
            except Exception as e:
                show_error(e)

# ----- Writing tools -----
TASKS = {
    # --- Basics ---
    "📝 Summarize": "Summarize the text in a few clear bullet points.",
    "✨ Rewrite professionally": "Rewrite the text in a polished, professional tone. Keep the meaning.",
    "🧒 Explain like I'm 12": "Explain the text in very simple words a 12-year-old would understand.",
    "✅ Fix grammar": "Fix spelling and grammar only. Return just the corrected text.",
    # --- Learn ---
    "🎓 Explain at expert level (with references)": (
        "Explain the topic or text at an advanced, expert level with precise terminology and depth. "
        "End with a 'Further reading' section suggesting well-known books, papers, or reputable "
        "organizations and websites the reader can search for. Do NOT invent links or citations; "
        "if you are not sure a source exists, say so."
    ),
    "🧠 Explain with analogies": "Explain the text using 3 vivid everyday analogies, then a one-line recap.",
    "📚 Make study notes": "Turn the text into well-organized study notes with headings, bullets and key terms in bold.",
    "🃏 Make flashcards": "Create 10 flashcards from the text in the format 'Q: ... / A: ...'.",
    "🧩 Make a quiz (5 questions)": (
        "Create a 5-question multiple-choice quiz from the text. Give 4 options per question, "
        "then an answer key with a one-line explanation for each."
    ),
    "🗺️ Mind-map outline": "Turn the text into a hierarchical mind-map outline using nested bullets.",
    "❓ Questions to think about": "Write 8 thought-provoking discussion questions based on the text, from easy to deep.",
    # --- Create ---
    "🤯 Interesting facts & trivia": (
        "Give 8 surprising, interesting facts or trivia related to the text's topic. "
        "Only include facts you are confident are true."
    ),
    "📖 Turn into a short story": "Turn the ideas in the text into an engaging short story of about 250 words.",
    "🎭 Turn into a poem": "Turn the text into a creative poem with a clear theme and rhythm.",
    "📱 Social media post": "Write a catchy social media post from the text with a hook, short lines, emojis and 5 hashtags.",
    "💼 LinkedIn post": "Write an engaging, professional LinkedIn post from the text with a strong first line and a closing question.",
    "📧 Draft an email": "Turn the text into a clear, polite email with a subject line.",
    "🎤 Presentation script": "Turn the text into a 2-minute spoken presentation script with an opening, 3 main points and a closing.",
    "🏷️ Title & headline ideas": "Suggest 10 creative titles/headlines for the text, mixing serious, catchy and curious styles.",
    # --- Think ---
    "🎯 Key takeaways & action items": "List the key takeaways, then a short checklist of concrete action items.",
    "⚖️ Pros and cons": "List the main pros and cons of the ideas in the text in two clear sections, then give a balanced verdict.",
    "🔍 Critique & counter-arguments": "Critically review the text: point out weak spots, missing evidence, and the strongest counter-arguments.",
    "🌍 Translate": "Translate the text into {lang}. Keep the tone and meaning. Return only the translation.",
}

with tools_tab:
    task = st.selectbox("What would you like to do?", list(TASKS.keys()))
    language = ""
    if task == "🌍 Translate":
        language = st.text_input("Translate into which language?", placeholder="e.g. Urdu, French, Arabic")
    if task.startswith("🎓"):
        st.caption("ℹ️ References are suggestions from the AI. Please verify them before relying on them.")
    source = st.radio("Input", ["✍️ Paste text", "📎 Upload file"], horizontal=True,
                      label_visibility="collapsed")
    text, uploaded = "", None
    if source == "✍️ Paste text":
        text = st.text_area("Paste your text here", height=200)
    else:
        uploaded = st.file_uploader("Upload a document or image", type=FILE_TYPES)
        st.caption(f"Supported: PDF, Word (.docx), TXT, Markdown, CSV, PNG, JPG · up to {MAX_MB} MB")

    if st.button("Run ✨"):
        content, problem = None, None
        if source == "✍️ Paste text":
            if text.strip():
                content = text
            else:
                problem = "Please paste some text first."
        elif uploaded is None:
            problem = "Please upload a file first."
        else:
            content, problem = prepare_file(uploaded)

        if not problem and task == "🌍 Translate" and not language.strip():
            problem = "Please type the language to translate into."

        if problem:
            st.warning(problem)
        else:
            instruction = TASKS[task].replace("{lang}", language.strip())
            with st.container(border=True):
                try:
                    run_with_indicator(content, instruction, "✨ Crafting your result")
                except Exception as e:
                    show_error(e)
